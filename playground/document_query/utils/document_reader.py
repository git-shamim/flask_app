import os
import io
from werkzeug.datastructures import FileStorage
from PyPDF2 import PdfReader
import docx
import openpyxl

def extract_text(file: FileStorage) -> str:
    filename = file.filename.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(file)
    elif filename.endswith(".docx"):
        return extract_docx_text(file)
    elif filename.endswith(".xlsx"):
        return extract_excel_text(file)
    elif filename.endswith(".txt"):
        return extract_txt_text(file)
    else:
        raise ValueError(f"Unsupported file type: {filename}")


def extract_pdf_text(file: FileStorage) -> str:
    try:
        reader = PdfReader(file.stream)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        raise RuntimeError(f"Error reading PDF file: {e}")


def extract_docx_text(file: FileStorage) -> str:
    try:
        doc = docx.Document(file.stream)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception as e:
        raise RuntimeError(f"Error reading DOCX file: {e}")


def extract_excel_text(file: FileStorage) -> str:
    try:
        wb = openpyxl.load_workbook(filename=io.BytesIO(file.read()), data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.append(" ".join([str(cell) for cell in row if cell is not None]))
        return "\n".join(text)
    except Exception as e:
        raise RuntimeError(f"Error reading Excel file: {e}")


def extract_txt_text(file: FileStorage) -> str:
    try:
        return file.stream.read().decode("utf-8")
    except Exception as e:
        raise RuntimeError(f"Error reading TXT file: {e}")
